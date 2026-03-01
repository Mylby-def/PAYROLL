from rest_framework import viewsets, status
from rest_framework.decorators import action, api_view, permission_classes
from django.views.decorators.csrf import ensure_csrf_cookie
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated, AllowAny
from rest_framework.authentication import SessionAuthentication
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.models import User
from django.db.models import Q, Sum
from django.http import HttpResponse
from django.utils import timezone
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from .models import (
    Teacher, Subject, Rate, Bonus,
    PayrollSheet, PayrollEntry
)
from .serializers import (
    TeacherSerializer, SubjectSerializer,
    RateSerializer, BonusSerializer,
    PayrollSheetSerializer, PayrollSheetListSerializer, PayrollEntrySerializer
)
from rest_framework import serializers


class UserSerializer(serializers.Serializer):
    id = serializers.IntegerField()
    username = serializers.CharField()
    email = serializers.EmailField(required=False)
    first_name = serializers.CharField(required=False)
    last_name = serializers.CharField(required=False)
    
    def to_representation(self, instance):
        return {
            'id': instance.id,
            'username': instance.username,
            'email': instance.email or '',
            'first_name': instance.first_name or '',
            'last_name': instance.last_name or '',
        }


class TeacherViewSet(viewsets.ModelViewSet):
    queryset = Teacher.objects.all()
    serializer_class = TeacherSerializer
    permission_classes = [IsAuthenticated]


class SubjectViewSet(viewsets.ModelViewSet):
    queryset = Subject.objects.all()
    serializer_class = SubjectSerializer
    permission_classes = [IsAuthenticated]


class RateViewSet(viewsets.ModelViewSet):
    queryset = Rate.objects.all()
    serializer_class = RateSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = super().get_queryset()
        subject_id = self.request.query_params.get('subject')
        date = self.request.query_params.get('date')
        rate_kind = self.request.query_params.get('rate_kind')

        if subject_id:
            queryset = queryset.filter(subject_id=subject_id)
        if date:
            queryset = queryset.filter(
                Q(effective_from__lte=date) &
                (Q(effective_to__isnull=True) | Q(effective_to__gte=date))
            )
        if rate_kind:
            queryset = queryset.filter(rate_kind=rate_kind)
        return queryset


class BonusViewSet(viewsets.ModelViewSet):
    queryset = Bonus.objects.all()
    serializer_class = BonusSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = super().get_queryset()
        teacher_id = self.request.query_params.get('teacher')
        period_start = self.request.query_params.get('period_start')
        period_end = self.request.query_params.get('period_end')
        
        if teacher_id:
            queryset = queryset.filter(teacher_id=teacher_id)
        if period_start:
            queryset = queryset.filter(period_end__gte=period_start)
        if period_end:
            queryset = queryset.filter(period_start__lte=period_end)
        
        return queryset


class PayrollSheetViewSet(viewsets.ModelViewSet):
    queryset = PayrollSheet.objects.all()
    permission_classes = [IsAuthenticated]

    def get_serializer_class(self):
        if self.action == 'list':
            return PayrollSheetListSerializer
        return PayrollSheetSerializer

    def get_queryset(self):
        queryset = super().get_queryset()
        status_filter = self.request.query_params.get('status')
        period_start = self.request.query_params.get('period_start')
        period_end = self.request.query_params.get('period_end')
        
        if status_filter:
            queryset = queryset.filter(status=status_filter)
        if period_start:
            queryset = queryset.filter(period_end__gte=period_start)
        if period_end:
            queryset = queryset.filter(period_start__lte=period_end)
        
        return queryset

    def perform_create(self, serializer):
        teacher = serializer.validated_data.get('teacher')
        period_start = serializer.validated_data.get('period_start')
        period_end = serializer.validated_data.get('period_end')
        title = serializer.validated_data.get('title') or ''
        if teacher and not title:
            title = f"Расчётный лист: {teacher.full_name} ({period_start} — {period_end})"
        serializer.save(created_by=self.request.user, title=title or 'Без названия')

    def destroy(self, request, *args, **kwargs):
        sheet = self.get_object()
        if sheet.status != 'draft':
            return Response(
                {'detail': 'Удалять можно только расчётные листы в статусе «Черновик».'},
                status=status.HTTP_400_BAD_REQUEST
            )
        return super().destroy(request, *args, **kwargs)

    def list(self, request, *args, **kwargs):
        response = super().list(request, *args, **kwargs)
        # Добавляем total_amount для каждого листа
        data_list = response.data.get('results', response.data)
        if not isinstance(data_list, list):
            data_list = [data_list] if data_list else []
        for item in data_list:
            if isinstance(item, dict) and 'id' in item:
                try:
                    sheet = PayrollSheet.objects.get(id=item['id'])
                    item['total_amount'] = str(sheet.get_total_amount())
                except PayrollSheet.DoesNotExist:
                    item['total_amount'] = '0.00'
        return response

    @action(detail=True, methods=['post'])
    def submit(self, request, pk=None):
        """Отправить расчётный лист"""
        sheet = self.get_object()
        sheet.status = 'submitted'
        sheet.save()
        return Response({'status': 'submitted'})

    @action(detail=True, methods=['post'])
    def approve(self, request, pk=None):
        """Утвердить расчётный лист"""
        sheet = self.get_object()
        sheet.status = 'approved'
        sheet.save()
        return Response({'status': 'approved'})

    @action(detail=True, methods=['get'])
    def export_excel(self, request, pk=None):
        """Экспорт расчётного листа в Excel"""
        sheet = self.get_object()
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Расчётный лист"
        
        # Стили
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        title_font = Font(bold=True, size=14)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Заголовок
        ws.merge_cells('A1:H1')
        ws['A1'] = f"Расчётный лист: {sheet.title}"
        ws['A1'].font = title_font
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        
        ws['A2'] = f"Период: {sheet.period_start} - {sheet.period_end}"
        ws['A3'] = f"Статус: {sheet.get_status_display()}"
        
        # Заголовки таблицы
        headers = ['№', 'Дата', 'Преподаватель', 'Предмет', 'Часы', 'Тариф за час', 'Тариф за группу', 'Сумма']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
        
        # Данные
        row = 6
        total = 0
        for idx, entry in enumerate(sheet.entries.all().order_by('date', 'teacher'), 1):
            ws.cell(row=row, column=1, value=idx).border = border
            ws.cell(row=row, column=2, value=entry.date.strftime('%Y-%m-%d')).border = border
            ws.cell(row=row, column=3, value=entry.teacher.full_name).border = border
            ws.cell(row=row, column=4, value=entry.subject.name).border = border
            ws.cell(row=row, column=5, value=float(entry.hours)).border = border
            ws.cell(row=row, column=6, value=float(entry.rate_per_hour) if entry.rate_per_hour else '').border = border
            ws.cell(row=row, column=7, value=float(entry.rate_per_group) if entry.rate_per_group else '').border = border
            ws.cell(row=row, column=8, value=float(entry.amount)).border = border
            total += entry.amount
            row += 1
        
        # Итого
        ws.cell(row=row, column=7, value="Итого:").font = Font(bold=True)
        ws.cell(row=row, column=8, value=float(total)).font = Font(bold=True)
        ws.cell(row=row, column=8).border = border
        
        # Ширина столбцов
        column_widths = [5, 12, 25, 20, 8, 12, 12, 12]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        
        # Сохранение
        http_response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f"payroll_{sheet.id}_{sheet.period_start}_{sheet.period_end}.xlsx"
        http_response['Content-Disposition'] = f'attachment; filename="{filename}"'
        wb.save(http_response)
        return http_response


@api_view(['GET'])
@permission_classes([AllowAny])
@ensure_csrf_cookie
def csrf_cookie_view(request):
    """Установить CSRF cookie для SPA (нужно вызвать до первого POST)."""
    return Response({'detail': 'CSRF cookie set'})


@api_view(['POST'])
@permission_classes([AllowAny])
def login_view(request):
    username = request.data.get('username')
    password = request.data.get('password')
    
    if not username or not password:
        return Response(
            {'detail': 'Необходимо указать имя пользователя и пароль'},
            status=status.HTTP_400_BAD_REQUEST
        )
    
    user = authenticate(request, username=username, password=password)
    if user is not None:
        login(request, user)
        serializer = UserSerializer(user)
        return Response({
            'user': serializer.data
        })
    else:
        return Response(
            {'detail': 'Неверные учетные данные'},
            status=status.HTTP_401_UNAUTHORIZED
        )


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def logout_view(request):
    logout(request)
    return Response({'detail': 'Выход выполнен успешно'})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def user_view(request):
    serializer = UserSerializer(request.user)
    return Response(serializer.data)


class PayrollEntryViewSet(viewsets.ModelViewSet):
    queryset = PayrollEntry.objects.all()
    serializer_class = PayrollEntrySerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        queryset = super().get_queryset()
        sheet_id = self.request.query_params.get('payroll_sheet')
        teacher_id = self.request.query_params.get('teacher')
        
        if sheet_id:
            queryset = queryset.filter(payroll_sheet_id=sheet_id)
        if teacher_id:
            queryset = queryset.filter(teacher_id=teacher_id)
        
        return queryset
